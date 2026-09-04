from pathlib import Path
from docx import Document


def load_docx(file_path):
    """
    Extract text, tables, and images from a DOCX document.
    """

    file_path = Path(file_path)

    document = Document(file_path)

    # -------------------------
    # Extract paragraphs
    # -------------------------

    paragraphs = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    # -------------------------
    # Extract tables
    # -------------------------

    tables = []

    for table_number, table in enumerate(document.tables, start=1):

        rows = []

        for row in table.rows:

            rows.append([
                cell.text.strip()
                for cell in row.cells
            ])

        tables.append({
            "table_number": table_number,
            "data": rows
        })

    # -------------------------
    # Extract images
    # -------------------------

    images = []

    image_dir = Path("data/extracted_docx_images")
    image_dir.mkdir(parents=True, exist_ok=True)

    for image_number, relationship in enumerate(
        document.part.rels.values(),
        start=1
    ):

        if "image" not in relationship.target_ref:
            continue

        image_part = relationship.target_part

        image_extension = Path(
            image_part.partname
        ).suffix

        image_filename = (
            f"image_{image_number}{image_extension}"
        )

        image_path = image_dir / image_filename

        with open(image_path, "wb") as image_file:
            image_file.write(image_part.blob)

        images.append({
            "image_number": image_number,
            "path": str(image_path)
        })

    return {
        "type": "docx",
        "paragraphs": paragraphs,
        "tables": tables,
        "images": images
    }