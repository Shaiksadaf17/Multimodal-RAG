from pathlib import Path
import pymupdf
from docx import Document
from pptx import Presentation

from src.extraction.pdf_tables import extract_tables_from_pdf
from src.extraction.pdf_images import extract_images_from_pdf


def load_pdf(file_path):
    """Extract text, tables, and images from a PDF document."""

    file_path = Path(file_path)

    # Open PDF using PyMuPDF
    document = pymupdf.open(file_path)

    pages = []

    for page_number, page in enumerate(document, start=1):
        text = page.get_text()

        pages.append({
            "page": page_number,
            "text": text
        })

    document.close()

    # Extract tables
    tables = extract_tables_from_pdf(file_path)

    # Extract images
    images = extract_images_from_pdf(file_path)

    return {
        "type": "pdf",
        "pages": pages,
        "tables": tables,
        "images": images
    }


def load_docx(file_path):
    """Extract text from a Word document."""

    file_path = Path(file_path)

    document = Document(file_path)

    paragraphs = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            paragraphs.append(paragraph.text.strip())

    return {
        "type": "docx",
        "paragraphs": paragraphs
    }


def load_pptx(file_path):
    """Extract text from a PowerPoint presentation."""

    file_path = Path(file_path)

    presentation = Presentation(file_path)

    slides = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        slide_text = []

        for shape in slide.shapes:

            if hasattr(shape, "text") and shape.text.strip():
                slide_text.append(shape.text.strip())

        slides.append({
            "slide": slide_number,
            "text": "\n".join(slide_text)
        })

    return {
        "type": "pptx",
        "slides": slides
    }


def load_document(file_path):
    """Automatically select the correct loader based on file type."""

    file_path = Path(file_path)

    extension = file_path.suffix.lower()

    if extension == ".pdf":
        return load_pdf(file_path)

    elif extension == ".docx":
        return load_docx(file_path)

    elif extension == ".pptx":
        return load_pptx(file_path)

    else:
        raise ValueError(
            f"Unsupported file type: {extension}"
        )