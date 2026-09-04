from pathlib import Path
from pptx import Presentation


def load_pptx(file_path):
    """
    Extract text, tables, and images from a PPTX presentation.
    """

    file_path = Path(file_path)
    presentation = Presentation(file_path)

    slides = []

    for slide_number, slide in enumerate(presentation.slides, start=1):

        slide_text = []
        slide_tables = []
        slide_images = []

        for shape in slide.shapes:

            # -------------------------
            # Extract text
            # -------------------------
            if hasattr(shape, "text") and shape.text.strip():
                slide_text.append(shape.text.strip())

            # -------------------------
            # Extract tables
            # -------------------------
            if shape.has_table:
                rows = []

                for row in shape.table.rows:
                    rows.append([
                        cell.text.strip()
                        for cell in row.cells
                    ])

                slide_tables.append(rows)

            # -------------------------
            # Extract images
            # -------------------------
            if shape.shape_type == 13:  # Picture

                image = shape.image

                image_dir = Path("data/extracted_pptx_images")
                image_dir.mkdir(
                    parents=True,
                    exist_ok=True
                )

                image_name = (
                    f"slide_{slide_number}_image_"
                    f"{len(slide_images) + 1}.{image.ext}"
                )

                image_path = image_dir / image_name

                with open(image_path, "wb") as image_file:
                    image_file.write(image.blob)

                slide_images.append({
                    "path": str(image_path)
                })

        slides.append({
            "slide": slide_number,
            "text": slide_text,
            "tables": slide_tables,
            "images": slide_images
        })

    return {
        "type": "pptx",
        "slides": slides
    }